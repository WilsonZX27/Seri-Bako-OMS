#!/usr/bin/env python3
"""Strict Seri Bako schedule workbook parser. Dry-run is the only default.

Catalog JSON is optional for structural inspection and required before a draft
payload can be considered ready. It has this shape:
  {"branches":[{"id":"...","name":"Gurney Plaza","status":"draft"}],
   "staff":[{"id":"...","name":"LEE SEE KIM","branch_id":"..."}],
   "current_assignments":[{...}]}

This tool never connects to Supabase. Its output is an auditable payload for
the authenticated import endpoint; remote writing remains a separate action.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

try:
    import openpyxl
except ModuleNotFoundError as error:
    raise SystemExit("openpyxl is required. Install scripts/attendance/requirements.txt first.") from error


DEFAULT_WEEK_START = dt.date(2026, 8, 31)
SHEET_BRANCHES = {
    "GURNEY PLAZA": "Gurney Plaza",
    "QUEENSBAY": "Queensbay Mall",
    "PUBLIKA": "Kwan Kee @ Publika",
    "IOI PUCHONG": "IOI Mall Puchong",
    "JAKEL SQUARE": "Jakel Mall",
    "TRX": "TRX",
    "PETALING STREET": "Kwan Kee Store @ Petaling Street",
    "SUNWAY PYRAMID": "Kwan Kee Store @ Sunway Pyramid",
    "MID VALLEY": "Kwan Kee Store @ Mid Valley",
    "VELOCITY": "Kwan Kee Store @ Sunway Velocity",
}
KINDS = {
    "OFF": "off", "OFF DAY": "off", "PH": "ph", "AL": "al",
    "MC": "mc", "UPL": "upl", "RO": "ro", "OIL": "oil",
}
BLOCKED_TEXT = {"CLEAR OT", "RESIGN", "LATE"}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def key(value: object) -> str:
    return clean(value).casefold()


def hhmm(value: object) -> str | None:
    if isinstance(value, dt.time):
        return value.strftime("%H:%M")
    if isinstance(value, dt.datetime):
        return value.strftime("%H:%M")
    return None


def day_for_header(value: object) -> int | None:
    if not isinstance(value, (int, float)):
        return None
    day = int(value)
    return day if 1 <= day <= 31 else None


def duration_minutes(start: str, end: str) -> int:
    start_minutes = int(start[:2]) * 60 + int(start[3:])
    end_minutes = int(end[:2]) * 60 + int(end[3:])
    return end_minutes - start_minutes if end_minutes > start_minutes else end_minutes + 1440 - start_minutes


def load_catalog(path: Path | None) -> dict:
    if path is None:
        return {"branches": [], "staff": [], "current_assignments": []}
    return json.loads(path.read_text(encoding="utf-8"))


def current_signature(row: dict) -> tuple:
    return (
        row.get("assignment_kind"), row.get("custom_start"), row.get("custom_end"),
        row.get("shift_template_id"), row.get("custom_break_seconds"),
        row.get("custom_maximum_recognised_seconds"),
    )


def parse(workbook_path: Path, catalog: dict, week_start: dt.date = DEFAULT_WEEK_START) -> dict:
    if week_start.weekday() != 0:
        raise ValueError("week_start must be a Monday")
    week_end = week_start + dt.timedelta(days=6)
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    branch_by_name = {key(row.get("name")): row for row in catalog.get("branches", [])}
    staff_by_name = defaultdict(list)
    for row in catalog.get("staff", []):
        staff_by_name[key(row.get("name"))].append(row)
    current = {
        (str(row.get("branch_id")), str(row.get("staff_id")), str(row.get("work_date"))): row
        for row in catalog.get("current_assignments", [])
    }
    unresolved: list[dict] = []
    cross_store: list[dict] = []
    suspicious_times: list[dict] = []
    outlet_reports: list[dict] = []
    draft_payloads: list[dict] = []

    missing_sheets = sorted(set(SHEET_BRANCHES) - set(workbook.sheetnames))
    extra_sheets = sorted(set(workbook.sheetnames) - set(SHEET_BRANCHES))
    if missing_sheets or extra_sheets:
        unresolved.append({"reason": "worksheet_set_mismatch", "missing": missing_sheets, "extra": extra_sheets})

    for sheet_name, canonical_name in SHEET_BRANCHES.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        branch_matches = [row for name, row in branch_by_name.items() if name == key(canonical_name)]
        branch = branch_matches[0] if len(branch_matches) == 1 else None
        if branch is None:
            unresolved.append({"sheet": sheet_name, "outlet": canonical_name, "reason": "canonical_outlet_not_found"})
        elif str(branch.get("status", "active")).casefold() in {"published", "locked"}:
            unresolved.append({"sheet": sheet_name, "outlet": canonical_name, "reason": "published_or_locked_week"})

        date_row = max(range(1, 4), key=lambda row: sum(isinstance(sheet.cell(row, column).value, (int, float)) for column in range(3, sheet.max_column + 1, 2)))
        all_date_columns = [(column, day_for_header(sheet.cell(date_row, column).value)) for column in range(3, sheet.max_column + 1, 2)]
        expected_dates = [week_start + dt.timedelta(days=index) for index in range(7)]
        date_columns: list[tuple[int, dt.date]] = []
        for index in range(0, len(all_date_columns) - 6):
            candidate = all_date_columns[index:index + 7]
            if [value for _, value in candidate] == [value.day for value in expected_dates]:
                date_columns = [(column, expected_dates[offset]) for offset, (column, _) in enumerate(candidate)]
                break
        if [work_date for _, work_date in date_columns] != expected_dates:
            unresolved.append({"sheet": sheet_name, "reason": "date_window_mismatch", "dates": [str(value) for _, value in date_columns]})

        assignments: list[dict] = []
        parsed_cells = 0
        sheet_unresolved_before = len(unresolved)
        for row_number in range(date_row + 2, sheet.max_row + 1):
            employee_name = clean(sheet.cell(row_number, 1).value)
            if not employee_name or key(employee_name) == "remark":
                break
            candidates = staff_by_name.get(key(employee_name), [])
            if branch:
                same_branch = [row for row in candidates if str(row.get("branch_id")) == str(branch.get("id"))]
            else:
                same_branch = []
            staff = same_branch[0] if len(same_branch) == 1 else None
            if staff is None:
                reason = "cross_outlet_employee" if candidates and not same_branch else "employee_not_found" if not candidates else "ambiguous_employee"
                issue = {"sheet": sheet_name, "row": row_number, "employee": employee_name, "reason": reason}
                unresolved.append(issue)
                if reason == "cross_outlet_employee":
                    cross_store.append(issue)

            for column, work_date in date_columns:
                first = sheet.cell(row_number, column).value
                second = sheet.cell(row_number, column + 1).value
                start, end = hhmm(first), hhmm(second)
                assignment: dict | None = None
                reason: str | None = None
                if start and end:
                    minutes = duration_minutes(start, end)
                    if minutes > 16 * 60:
                        reason = "suspicious_time_span"
                        suspicious_times.append({"sheet": sheet_name, "row": row_number, "employee": employee_name, "date": str(work_date), "start": start, "end": end, "minutes": minutes})
                    else:
                        assignment = {"assignment_kind": "shift", "shift_template_id": None, "custom_start": start, "custom_end": end, "custom_crosses_midnight": end <= start, "custom_reason": "Imported from September 2026 workbook"}
                elif start or end:
                    reason = "incomplete_time_pair"
                else:
                    label = clean(first).upper()
                    if not label and not clean(second):
                        continue
                    if label in KINDS and not clean(second):
                        assignment = {"assignment_kind": KINDS[label]}
                    elif label in BLOCKED_TEXT:
                        reason = "blocked_instruction"
                    else:
                        reason = "unsupported_or_cross_store_value"
                if reason:
                    unresolved.append({"sheet": sheet_name, "row": row_number, "employee": employee_name, "date": str(work_date), "value": clean(first), "value_2": clean(second), "reason": reason})
                elif assignment:
                    parsed_cells += 1
                    if branch and staff:
                        assignment.update({"staff_id": staff["id"], "work_date": str(work_date)})
                        assignments.append(assignment)

        differences = []
        if branch:
            proposed = {(str(branch["id"]), str(row["staff_id"]), row["work_date"]): row for row in assignments}
            relevant_current = {k: v for k, v in current.items() if k[0] == str(branch["id"]) and week_start.isoformat() <= k[2] <= week_end.isoformat()}
            for assignment_key in sorted(set(proposed) | set(relevant_current)):
                before, after = relevant_current.get(assignment_key), proposed.get(assignment_key)
                if before is None or after is None or current_signature(before) != current_signature(after):
                    differences.append({"staff_id": assignment_key[1], "work_date": assignment_key[2], "before": before, "after": after})
        ready = bool(branch) and len(unresolved) == sheet_unresolved_before and len(assignments) > 0
        outlet_reports.append({"sheet": sheet_name, "outlet": canonical_name, "branch_id": branch.get("id") if branch else None, "parsed_cell_count": parsed_cells, "matched_assignment_count": len(assignments), "diff_count": len(differences), "ready": ready})
        if branch:
            draft_payloads.append({"branch_id": branch["id"], "outlet": canonical_name, "assignments": assignments, "differences": differences})

    parsed_cells = sum(row["parsed_cell_count"] for row in outlet_reports)
    total_cells = parsed_cells + len([row for row in unresolved if row.get("date")])
    matched_cells = sum(row["matched_assignment_count"] for row in outlet_reports)
    return {
        "mode": "dry-run",
        "source_file": workbook_path.name,
        "source_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "week_start": str(week_start),
        "week_end": str(week_end),
        "worksheets": len(workbook.sheetnames),
        "outlets": outlet_reports,
        "totals": {"parseable_cells": parsed_cells, "matched_assignments": matched_cells, "candidate_cells": total_cells, "match_rate": round(matched_cells / total_cells, 4) if total_cells else 0, "unresolved": len(unresolved)},
        "unresolved_by_reason": dict(Counter(row["reason"] for row in unresolved)),
        "unresolved": unresolved,
        "suspicious_times": suspicious_times,
        "cross_store_records": cross_store,
        "draft_payloads": draft_payloads,
        "ready_for_authenticated_draft_import": len(outlet_reports) == 10 and all(row["ready"] for row in outlet_reports),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Dry-run one Monday-Sunday window from an OMS schedule workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--week-start", type=dt.date.fromisoformat, default=DEFAULT_WEEK_START, help="Monday in YYYY-MM-DD format (default: 2026-08-31)")
    parser.add_argument("--catalog", type=Path, help="Canonical branch/staff/current-assignment JSON export")
    parser.add_argument("--output", type=Path, help="Write the full JSON report")
    args = parser.parse_args()
    if not args.workbook.is_file():
        parser.error(f"Workbook not found: {args.workbook}")
    try:
        report = parse(args.workbook, load_catalog(args.catalog), args.week_start)
    except ValueError as error:
        parser.error(str(error))
    encoded = json.dumps(report, ensure_ascii=False, indent=2, default=str)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(encoded + "\n", encoding="utf-8")
    print(json.dumps({key: report[key] for key in ("mode", "source_sha256", "week_start", "week_end", "worksheets", "totals", "unresolved_by_reason", "ready_for_authenticated_draft_import")}, ensure_ascii=False, indent=2))
    return 0 if report["ready_for_authenticated_draft_import"] else 2


if __name__ == "__main__":
    sys.exit(main())
