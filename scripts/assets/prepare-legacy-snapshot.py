"""Read-only asset export preparation. Never imports accounts or credentials.

Use the bundled Python runtime with openpyxl. Output stays outside Git.
"""
import argparse
import datetime
import hashlib
import json
from pathlib import Path

import openpyxl

TABLE_KEYS = {
    "Asset Master List": "Asset ID",
    "Maintenance Requests": "Request ID",
    "Maintenance Updates": "Update ID",
    "Asset Transfer Requests": "Transfer ID",
    "Asset Transfer Updates": "Update ID",
    "Daily Audit Logs": None,  # One submission may contain many asset rows.
    "Attachments": None,
    "Suppliers": "Supplier ID",
}


def value(cell):
    if isinstance(cell, (datetime.datetime, datetime.date)):
        return cell.isoformat()
    return cell


def prepare(source):
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    tables, issues, warnings = {}, [], []
    for name, key in TABLE_KEYS.items():
        if name not in workbook.sheetnames:
            issues.append({"table": name, "reason": "Missing source sheet"})
            continue
        iterator = workbook[name].iter_rows(values_only=True)
        headers = next(iterator)
        populated = [(i, str(h)) for i, h in enumerate(headers) if h is not None]
        if len({h for _, h in populated}) != len(populated):
            issues.append({"table": name, "reason": "Duplicate column names"})
        rows, seen = [], set()
        for number, cells in enumerate(iterator, 2):
            if not any(c is not None for c in cells):
                continue
            row = {h: value(cells[i]) if i < len(cells) else None for i, h in populated}
            row["_source_row"] = number
            # Row keys keep repeated submissions / updates distinct and repeatable.
            row["_source_key"] = f"{name}:{number}"
            if key:
                identity = str(row.get(key) or "").strip().upper()
                if not identity or identity in seen:
                    issues.append({"table": name, "row": number, "reason": "Missing or duplicate ID", "id": identity})
                seen.add(identity)
            if any(cells[i] is not None for i, h in enumerate(headers) if h is None and i < len(cells)):
                issues.append({"table": name, "row": number, "reason": "Nonempty column without header"})
            rows.append(row)
        tables[name] = rows
    assets = tables.get("Asset Master List", [])
    asset_ids = {str(r.get("Asset ID") or "").strip().upper() for r in assets}
    request_ids = {str(r.get("Request ID") or "").strip() for r in tables.get("Maintenance Requests", [])}
    transfer_ids = {str(r.get("Transfer ID") or "").strip() for r in tables.get("Asset Transfer Requests", [])}
    quantities = []
    for row in assets:
        quantity = row.get("Qty")
        if quantity is None:
            warnings.append({"table": "Asset Master List", "row": row["_source_row"], "reason": "Source quantity is blank; preserve unknown holding"})
        elif isinstance(quantity, bool) or not isinstance(quantity, (int, float)) or quantity < 0 or quantity != int(quantity):
            issues.append({"table": "Asset Master List", "row": row["_source_row"], "reason": "Invalid quantity"})
        else:
            quantities.append(quantity)
    for table in ["Maintenance Requests", "Asset Transfer Requests", "Daily Audit Logs"]:
        for row in tables.get(table, []):
            identity = str(row.get("Asset ID") or "").strip().upper()
            if identity not in asset_ids:
                issues.append({"table": table, "row": row["_source_row"], "reason": "Unmatched asset", "id": identity})
    for table, key, ids in [("Maintenance Updates", "Request ID", request_ids), ("Asset Transfer Updates", "Transfer ID", transfer_ids)]:
        for row in tables.get(table, []):
            identity = str(row.get(key) or "").strip()
            if identity not in ids:
                issues.append({"table": table, "row": row["_source_row"], "reason": "Unmatched parent", "id": identity})
    attachments = []
    for table, rows in tables.items():
        for row in rows:
            for field, content in row.items():
                if ("URL" in field or field == "File URL") and content:
                    attachments.append({"table": table, "row": row["_source_row"], "field": field, "url": content, "verified": False})
    summary = {"rows": {name: len(rows) for name, rows in tables.items()}, "unique_asset_ids": len(asset_ids), "quantity_total": sum(quantities), "multi_quantity_assets": sum(q > 1 for q in quantities), "locations": sorted({str(r.get("Location") or "") for r in assets}), "validation_issues": len(issues), "unknown_quantities": len(warnings), "unresolved_relations": sum(i["reason"].startswith("Unmatched") for i in issues), "attachments_to_verify": len(attachments), "ready_for_cutover": False}
    return {"source_sha256": hashlib.sha256(Path(source).read_bytes()).hexdigest(), "tables": tables, "issues": issues, "warnings": warnings, "attachments": attachments, "summary": summary}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    snapshot = prepare(args.source)
    args.output.mkdir(parents=True, exist_ok=True)
    (args.output / "snapshot.json").write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.output / "reconciliation.json").write_text(json.dumps({"summary": snapshot["summary"], "issues": snapshot["issues"], "warnings": snapshot["warnings"]}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(snapshot["summary"], ensure_ascii=False))
